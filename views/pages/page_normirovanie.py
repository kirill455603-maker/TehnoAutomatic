
from ui_page_normirovanie import Ui_PageNormirovanie
from datetime import datetime
from PySide6.QtWidgets import (
    QWidget, QPushButton, QFileDialog, QHBoxLayout, QVBoxLayout,
    QLabel, QGroupBox, QButtonGroup, QComboBox, QCompleter,
    QInputDialog, QCheckBox, QLineEdit, QFrame, QScrollArea
)
from PySide6.QtCore import Signal, Qt, QTimer
from openpyxl import load_workbook

from views.widgets.widgets_normirovanie.w_normirovanie_button import ButtonList


class PageNormirovanie(QWidget):
    btn_home_clicked = Signal()
    editing_combobox = Signal()
    btn_result_clicked = Signal()
    btn_show_operations_clicked = Signal()
    def __init__(self, parent=None):
        super().__init__(parent)
        self.ui = Ui_PageNormirovanie()
        self.ui.setupUi(self)
        self.create_completer(self.ui.comboBox_details)
        self.create_completer(self.ui.comboBox_operations)
        self.hide_widgets()

        self.path_file = None
        self.wb = None
        self.selected_sheet_name = None
        self.sheet = None

        self.details_list = []
        self.details_map = {}
        self.operations_list = []
        self.operations_map = {}

        self.buttons_list = []
        self.buttons_list_operation = []
        self.widgets_operations = []

        self.ui.btn_home.clicked.connect(self._on_button_home)
        self.ui.btn_open_file.clicked.connect(self._on_btn_open_file)
        self.ui.btn_clear_details.clicked.connect(self._clear_details)
        self.ui.btn_clear_operations.clicked.connect(self._clear_operations)
        self.ui.pushButton_calc_result.clicked.connect(
            lambda: self.btn_result_clicked.emit()
        )
        self.ui.pushButton_show_btn_operations.clicked.connect(self._on_btn_show_operations)

        self.ui.comboBox_operations.currentTextChanged.connect(self._combobox_edit)

    def add_log(self, text, type_msg = 'info',):
        now = datetime.now()
        time = now.strftime('%d.%m.%Y %H:%M:%S')
        msg = ''
        if type_msg == 'info':
            msg = f'[{time}] [ИНФОРМАЦИЯ]: {text}'
        elif type_msg == 'error':
            msg = f'[{time}] [ОШИБКА]: {text}'
        self.ui.plainTextEdit_log.appendPlainText(msg)
        scrollbar = self.ui.plainTextEdit_log.verticalScrollBar()
        scrollbar.setValue(scrollbar.maximum())

    def _on_btn_show_operations(self):
        self._clear_buttons_operations_tehno_pro()
        self.btn_show_operations_clicked.emit()

    def get_row_and_col(self):
        find_detail = self.ui.comboBox_details.currentText()
        find_operation = self.ui.comboBox_operations.currentText()

        row = None
        col = None
        for i, detail in enumerate(self.details_list):
            if detail == find_detail:
                row = i + 2
                break
        for i, operation in enumerate(self.operations_list):
            if operation == find_operation:
                col = i + 3
                break
        return row, col

    def auto_scroll_down(self):
        scroll = self.ui.scrollArea.verticalScrollBar()
        scroll.setValue(scroll.maximum())

    def get_value(self):
        return self.ui.comboBox_details.currentText(), self.ui.comboBox_operations.currentText()

    def _combobox_edit(self):
        #Очищаем перед сменой
        self.clear_widgets()
        self.editing_combobox.emit()

    def _clear_details(self):
        self.ui.comboBox_details.setCurrentText('')

    def _clear_operations(self):
        self.ui.comboBox_operations.setCurrentText('')

    def create_completer(self, widget):
        completer = QCompleter(widget.model())
        completer.setCaseSensitivity(Qt.CaseInsensitive)
        completer.setFilterMode(Qt.MatchContains)
        completer.setCompletionMode(QCompleter.PopupCompletion)
        widget.setCompleter(completer)

    def hide_widgets(self):
        self.ui.find_operations.hide()
        self.ui.widget_with_details.hide()
        self.ui.widget_with_buttons.hide()
        self.ui.plainTextEdit_log.hide()

    def add_widgets(self, widget):
        self.ui.layout_widgets.addWidget(widget)
        self.widgets_operations.append(widget)

    def clear_widgets(self):
        for widget in self.widgets_operations:
            widget.setParent(None)
            widget.deleteLater()
        self.widgets_operations.clear()

    def save_result_in_excel(self, row, col, time):
        if row is None or col is None:
            self.show_text('❌ Не удалось определить координаты ячейки')
            raise ValueError("Не удалось определить координаты ячейки")

        self.sheet.cell(row=row, column=col, value=time)

        # ✅ Сохраняем файл на диск
        try:
            self.wb.save(self.path_file)
        except PermissionError:
            self.show_text('❌ Не удалось сохранить файл. Возможно, он открыт в Excel.')
            raise PermissionError(
                "Не удалось сохранить файл. Возможно, он открыт в Excel."
            )

    def _on_button_home(self):
        self.btn_home_clicked.emit()

    def _on_btn_open_file(self):
        #Очистка до чтения файла
        self.hide_widgets()
        self.ui.error_path.setText('')
        self._clear_buttons()

        file_path, _ = QFileDialog.getOpenFileName(
            self, 'Выберите EXCEL файл',
            r'\\mgsrv\LGMGroup\OGT\Персоналии\Гомов\технолоджия',
            'Excel files (*.xlsx)'
        )

        if not file_path:
            return

        if not file_path.lower().endswith(('.xlsx', '.xls')):
            self.ui.error_path.setText('Ошибка: Это не Excel файл!')
            return

        if self.wb is not None:
            self.wb.close()

        self.path_file = file_path
        self.wb = load_workbook(file_path)
        if len(self.wb.sheetnames) > 1:
            for sheet_name in self.wb.sheetnames:
                self._add_list_buttons(sheet_name)
                self.ui.widget_with_buttons.show()
        else:
            self.ui.widget_with_details.show()
            self.ui.find_operations.show()
            self.selected_sheet_name = self.wb.active
            self.sheet = self.wb.active
            self._add_details()
            self._add_operations()
        self.ui.label.setText(f'Открыт файл: {file_path}')
        self.add_log(f'Был открыт файл: {file_path}')
        self.ui.btn_open_file.setText('Выбрать другой файл')
        self.ui.error_path.setText('')

    def add_button_in_layout(self, num, text):
        btn = ButtonList()
        self.buttons_list_operation.append(btn)
        btn.ui.pushButton.clicked.connect(
            lambda: self._add_text_in_input_operations(text)
        )
        btn.edit_text(f'{num}. {text}')
        self.ui.layout_buttons_tehno_pro.addWidget(btn)

    def _add_text_in_input_operations(self, text):
        self.ui.comboBox_operations.setCurrentText(text)

    def _add_list_buttons(self, sheet_name):
        btn = ButtonList()
        btn.edit_text(sheet_name)
        btn.ui.pushButton.clicked.connect(
            lambda: self._select_sheet(btn.get_text_button()))
        self.ui.layout_with_buttons.addWidget(btn)
        self.buttons_list.append(btn)

    def _select_sheet(self, name):
        self.selected_sheet_name = name
        self.sheet = self.wb[name]
        self._add_details()
        self._add_operations()
        if self.details_list and self.operations_list:
            self.ui.widget_with_details.show()
            self.ui.find_operations.show()
        self._clear_buttons()
        self.ui.widget_with_buttons.hide()

    def _add_details(self):
        self.ui.comboBox_details.clear()
        self.details_list.clear()
        for i, row in enumerate(self.sheet.iter_rows(min_col=1, min_row=2, max_col=2, values_only=True)):
            if not row[0] or not row[1]:
                continue
            if row[0] == '*':
                detail = row[1]
            else:
                detail = f'{row[0]} {row[1]}'
            self.details_list.append(detail)

        if not self.details_list:
            self.ui.error_path.setText('Ошибка: Нет деталей для добавления. Убедитесь, что детали начинаются со второй строки')
            return

        self.ui.comboBox_details.addItems(self.details_list)
        self.ui.comboBox_details.setCurrentText('')
        self.ui.comboBox_details.lineEdit().setPlaceholderText('Введите название/номер детали или выберите из списка...')

    def _add_operations(self):
        self.ui.comboBox_operations.clear()
        self.operations_list.clear()
        for row in self.sheet.iter_rows(max_row=1, min_col=3, values_only=True):
            for el in row:
                if not el:
                    continue
                if el.lower() in ['общий итог', 'количество', 'итого на насос']:
                    continue
                self.operations_list.append(el)

        if not self.details_list:
            self.ui.error_path.setText('Ошибка: Нет операций для добавления. Убедитесь, что операции находятся на первой строке')
            return

        self.ui.comboBox_operations.addItems(self.operations_list)
        self.ui.comboBox_operations.setCurrentText('')
        self.ui.comboBox_operations.lineEdit().setPlaceholderText('Введите операцию или выберите из списка...')

    def _clear_buttons(self):
        for button in self.buttons_list:
            button.setParent(None)
            button.deleteLater()
        self.buttons_list.clear()

    def _clear_buttons_operations_tehno_pro(self):
        for button in self.buttons_list_operation:
            button.setParent(None)
            button.deleteLater()
        self.buttons_list_operation.clear()

    def show_text(self, text):
        self.ui.error_path.setText(text)



